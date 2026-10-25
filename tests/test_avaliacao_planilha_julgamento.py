"""Testes da planilha de julgamento (store e documentos falsos; sem banco)."""
import uuid

import openpyxl
import pytest

from avaliacao import planilha_julgamento as pj
from avaliacao.store import Julgamento, Par, Tese
from scraper.schema import DocumentoJuridico


class _FakeStore:
    def __init__(self, teses, pares):
        self.teses = teses
        self.pares = pares  # tese_id -> list[Par]
        self.julgamentos: dict[tuple, Julgamento] = {}

    def listar_teses(self, split=None):
        return list(self.teses)

    def pares_da_tese(self, tese_id):
        return list(self.pares.get(tese_id, []))

    def gravar_julgamento(self, j):
        self.julgamentos[(j.par_id, j.juiz, j.rodada)] = j


def _doc(i, ementa="Ementa do acórdão %d"):
    return DocumentoJuridico(
        id_documento=f"d{i}", tipo_documento="ACORDAO", hierarquia_categoria=1,
        numero_processo=f"0000{i}-00.2020.5.01.0001", ementa=ementa % i if "%d" in ementa else ementa,
        fundamentacao="F" * 4000, acordao="A" * 100 + "FIM", link_original=f"http://x/{i}",
    )


@pytest.fixture
def cenario():
    tese = Tese(uuid.uuid4(), {"pedido_principal": "adicional", "tese_central": "tese X",
                               "normas": ["CLT 192"], "_texto_original": "texto do advogado"},
                "advogado", "insalubridade", "dev", "d0")
    pares = [Par(uuid.uuid4(), tese.id, f"d{i}", ["hibrido", "bm25_top10"]) for i in range(12)]
    docs = {f"d{i}": _doc(i) for i in range(12)}
    docs["d3"] = _doc(3, ementa="")
    store = _FakeStore([tese], {tese.id: pares})
    return tese, pares, docs, store


def _avaliacao(caminho):
    return openpyxl.load_workbook(caminho)["Avaliação"]


def _linhas(ws):
    return [r for r in ws.iter_rows(min_row=pj.LINHA_INICIAL) if r[0].value]


def _preencher(caminho, notas, email="  Ana@Exemplo.COM ", stances=None):
    wb = openpyxl.load_workbook(caminho)
    wb["Avaliador"]["B2"] = email
    ws = wb["Avaliação"]
    for i, r in enumerate(_linhas(ws)):
        if i in notas:
            ws.cell(r[0].row, pj.COL_NOTA).value = notas[i]
        if stances and i in stances:
            ws.cell(r[0].row, pj.COL_STANCE).value = stances[i]
    wb.save(caminho)
    return [r[0].value for r in _linhas(ws)]


def test_exportar_tem_as_quatro_abas_e_linhas_sorteadas(cenario, tmp_path):
    tese, pares, docs, _ = cenario
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    assert pj.exportar_lote(tese, pares, docs, a, seed=1) == 12
    pj.exportar_lote(tese, pares, docs, b, seed=1)
    assert openpyxl.load_workbook(a).sheetnames == ["Instruções", "Rubrica", "Avaliador", "Avaliação"]
    ws = _avaliacao(a)
    ordem_a = [r[0].value for r in _linhas(ws)]
    assert ordem_a == [r[0].value for r in _linhas(_avaliacao(b))]
    assert ordem_a != [str(p.id) for p in sorted(pares, key=lambda p: p.documento_id)]
    c = tmp_path / "c.xlsx"
    pj.exportar_lote(tese, pares, docs, c, seed=2)
    assert ordem_a != [r[0].value for r in _linhas(_avaliacao(c))]
    # nenhuma pista de sistema/corredor
    wb = openpyxl.load_workbook(a)
    for aba in wb:
        for row in aba.iter_rows():
            for cel in row:
                assert "hibrido" not in str(cel.value or "").lower()
                assert "corredor" not in str(cel.value or "").lower()
    assert ws.freeze_panes == f"A{pj.LINHA_INICIAL}"
    assert ws.column_dimensions["A"].hidden
    # ementa vazia -> dispositivo (fim do acórdão)
    textos = [r[pj.COL_EMENTA - 1].value for r in _linhas(ws)]
    assert any(t.endswith("FIM") for t in textos)
    assert "texto do advogado" in " ".join(str(c.value) for row in ws.iter_rows(max_row=pj.LINHA_INICIAL - 1) for c in row if c.value)


def test_round_trip(cenario, tmp_path):
    tese, pares, docs, store = cenario
    caminho = tmp_path / "lote.xlsx"
    pj.exportar_lote(tese, pares, docs, caminho, seed=1)
    ids = _preencher(caminho, {0: 3, 1: 1, 2: 2}, stances={0: "fortalece", 1: "enfraquece", 2: "neutro"})
    email, linhas = pj.ler_lote(caminho)
    assert email == "ana@exemplo.com"
    assert {(str(p), n, s) for p, n, s in linhas} == {
        (ids[0], 3, "fortalece"), (ids[1], 1, "neutro"), (ids[2], 2, "neutro")}
    assert pj.importar_arquivo(store, caminho) == (3, {"sem nota": 9})
    j = store.julgamentos[(uuid.UUID(ids[0]), "humano:ana@exemplo.com", 1)]
    assert (j.relevancia, j.stance) == (3, "fortalece")


def test_importar_ignora_linhas_invalidas(cenario, tmp_path, capsys):
    tese, pares, docs, store = cenario
    caminho = tmp_path / "lote.xlsx"
    pj.exportar_lote(tese, pares, docs, caminho, seed=1)
    ids = _preencher(caminho, {0: 5, 1: 2, 2: "x", 3: 3, 4: 2},
                     stances={1: None, 3: "fortalece", 4: "inventado"})
    wb = openpyxl.load_workbook(caminho)
    ws = wb["Avaliação"]
    ws.cell(pj.LINHA_INICIAL + 5, 1).value = str(uuid.uuid4())  # par de outra tese
    ws.cell(pj.LINHA_INICIAL + 5, pj.COL_NOTA).value = 3
    ws.cell(pj.LINHA_INICIAL + 5, pj.COL_STANCE).value = "neutro"
    wb.save(caminho)
    gravadas, puladas = pj.importar_arquivo(store, caminho)
    assert gravadas == 1
    assert [k[0] for k in store.julgamentos] == [uuid.UUID(ids[3])]
    assert puladas["par fora da tese"] == 1
    assert sum(v for k, v in puladas.items() if k != "sem nota") == 5
    assert "linha" in capsys.readouterr().err


def test_importar_idempotente(cenario, tmp_path):
    tese, pares, docs, store = cenario
    caminho = tmp_path / "lote.xlsx"
    pj.exportar_lote(tese, pares, docs, caminho, seed=1)
    _preencher(caminho, {0: 3, 1: 0}, stances={0: "fortalece"})
    pj.importar_arquivo(store, caminho)
    antes = dict(store.julgamentos)
    pj.importar_arquivo(store, caminho)
    assert store.julgamentos == antes and len(antes) == 2


@pytest.mark.parametrize("email", [None, "   ", "semarroba"])
def test_email_ausente_falha(cenario, tmp_path, email):
    tese, pares, docs, _ = cenario
    caminho = tmp_path / "lote.xlsx"
    pj.exportar_lote(tese, pares, docs, caminho, seed=1)
    _preencher(caminho, {0: 3}, email=email, stances={0: "neutro"})
    with pytest.raises(ValueError, match="e-mail"):
        pj.ler_lote(caminho)


def test_sem_nenhuma_nota_falha(cenario, tmp_path):
    tese, pares, docs, _ = cenario
    caminho = tmp_path / "lote.xlsx"
    pj.exportar_lote(tese, pares, docs, caminho, seed=1)
    _preencher(caminho, {})
    with pytest.raises(ValueError, match="nota"):
        pj.ler_lote(caminho)


def test_nome_do_arquivo_e_cli_importar(cenario, tmp_path, capsys):
    tese, pares, docs, store = cenario
    assert pj.nome_arquivo(tese, "2026-10-22") == f"lote_2026-10-22_insalubridade_{str(tese.id)[:8]}.xlsx"
    ok = tmp_path / "ok.xlsx"
    pj.exportar_lote(tese, pares, docs, ok, seed=1)
    _preencher(ok, {0: 0})
    ruim = tmp_path / "ruim.xlsx"
    pj.exportar_lote(tese, pares, docs, ruim, seed=1)
    assert pj.main(["importar", str(ok)], store=store) == 0
    assert pj.main(["importar", str(ok), str(ruim)], store=store) == 1
    out = capsys.readouterr()
    assert "1 gravada" in out.out


def test_stance_ignorado_com_nota_baixa(cenario, tmp_path):
    tese, pares, docs, store = cenario
    caminho = tmp_path / "lote.xlsx"
    pj.exportar_lote(tese, pares, docs, caminho, seed=1)
    ids = _preencher(caminho, {0: 0, 1: 1}, stances={0: "inventado", 1: "fortalece"})
    email, linhas = pj.ler_lote(caminho)
    assert {(str(p), n, s) for p, n, s in linhas} == {(ids[0], 0, "neutro"), (ids[1], 1, "neutro")}


def test_cegueira_so_uuids(cenario, tmp_path):
    tese, pares, docs, _ = cenario
    caminho = tmp_path / "lote.xlsx"
    pj.exportar_lote(tese, pares, docs, caminho, seed=1)
    wb = openpyxl.load_workbook(caminho)
    assert str(uuid.UUID(wb.defined_names["tese_id"].attr_text.strip('"'))) == str(tese.id)
    ws = wb["Avaliação"]
    assert ws.column_dimensions["A"].hidden
    for r in _linhas(ws):
        uuid.UUID(r[0].value)


def test_cli_exportar_filtra_corredores_e_pula_tese_vazia(cenario, tmp_path, monkeypatch, capsys):
    tese, pares, docs, store = cenario
    vazia = Tese(uuid.uuid4(), {"tese_central": "y"}, "sintetica", "outro", "dev")
    antigo = Par(uuid.uuid4(), tese.id, "d99", ["bm25"])
    store.teses.append(vazia)
    store.pares[tese.id] = pares + [antigo]
    store.pares[vazia.id] = [Par(uuid.uuid4(), vazia.id, "d1", ["bm25", "pipeline"])]
    buscados = []

    class _Ret:
        def __init__(self, url): pass
        def fetch(self, ids):
            buscados.extend(ids)
            return docs

    monkeypatch.setenv("DATABASE_URL", "x")
    monkeypatch.setattr("query.search.doc_retriever.DocRetriever", _Ret)
    assert pj.main(["exportar", "--todas", "--pasta", str(tmp_path / "l")], store=store) == 0
    assert "d99" not in buscados
    arquivos = list((tmp_path / "l").glob("*.xlsx"))
    assert len(arquivos) == 1
    assert len(_linhas(_avaliacao(arquivos[0]))) == 12
    assert "sem pares do pool" in capsys.readouterr().err


def _lote_preenchido(cenario, tmp_path, email):
    tese, pares, docs, store = cenario
    caminho = tmp_path / "lote.xlsx"
    pj.exportar_lote(tese, pares, docs, caminho, seed=1)
    ids = _preencher(caminho, {0: 3}, email=email, stances={0: "fortalece"})
    return store, caminho, ids


def test_importar_adjudicacao_grava_rodada_2(cenario, tmp_path):
    store, caminho, ids = _lote_preenchido(cenario, tmp_path, "Alexia@x.com")
    assert pj.importar_arquivo(store, caminho, rodada=2, adjudicador=" alexia@X.com ")[0] == 1
    assert list(store.julgamentos) == [(uuid.UUID(ids[0]), "humano:alexia@x.com", 2)]


def test_importar_adjudicacao_com_outro_email_falha(cenario, tmp_path):
    store, caminho, _ = _lote_preenchido(cenario, tmp_path, "ana@x.com")
    with pytest.raises(ValueError, match="AVALIACAO_ADJUDICADOR"):
        pj.importar_arquivo(store, caminho, rodada=2, adjudicador="alexia@x.com")
    with pytest.raises(ValueError, match="AVALIACAO_ADJUDICADOR"):
        pj.importar_arquivo(store, caminho, rodada=2, adjudicador=None)
    assert store.julgamentos == {}


def test_cli_adjudicacao(cenario, tmp_path, monkeypatch, capsys):
    store, caminho, ids = _lote_preenchido(cenario, tmp_path, "alexia@x.com")
    monkeypatch.setenv("AVALIACAO_ADJUDICADOR", "alexia@x.com")
    assert pj.main(["importar", "--adjudicacao", str(caminho)], store=store) == 0
    assert [k[2] for k in store.julgamentos] == [2]
    monkeypatch.setenv("AVALIACAO_ADJUDICADOR", "outra@x.com")
    assert pj.main(["importar", "--adjudicacao", str(caminho)], store=store) == 1
    assert "AVALIACAO_ADJUDICADOR" in capsys.readouterr().err
