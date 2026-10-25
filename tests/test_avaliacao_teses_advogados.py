"""Importação das teses dos advogados (planilha preenchida), sem rede nem Ollama."""
import json
import os
import uuid

import openpyxl
import psycopg
import pytest

from avaliacao import teses_advogados as ta
from avaliacao.store import Tese
from avaliacao.temas import TEMAS
from query.enrichment.schema import EstruturaArgumentativa

LADO_RECLAMANTE = "A favor do reclamante (trabalhador)"
LADO_RECLAMADA = "A favor da reclamada (empresa)"


def _planilha(caminho, linhas, email="adv@x.com"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Teses"
    ws["A1"] = "Produção de teses — JuriSight"
    ws["A3"], ws["B3"] = "Nome", "Fulana"
    ws["A4"], ws["B4"] = "E-mail", email
    ws["A5"], ws["B5"] = "Atuação", "Reclamante"
    ws["A6"], ws["B6"] = "Anos de experiência em direito do trabalho", 5
    for col, t in zip("ABCDEFG", ["Nº", "Tema", "Processo", "Ementa do acórdão",
                                   "Tese — o que você digitaria no JuriSight",
                                   "A favor de quem é a tese", "Observações — opcional"]):
        ws[f"{col}8"] = t
    ws["A9"], ws["E9"], ws["F9"] = "Ex.", "tese de exemplo", LADO_RECLAMANTE
    for i, (n, tema, texto, lado) in enumerate(linhas, start=10):
        ws[f"A{i}"], ws[f"B{i}"], ws[f"E{i}"], ws[f"F{i}"] = n, tema, texto, lado
    wb.save(caminho)
    return caminho


class _Enricher:
    def __init__(self):
        self.chamadas = []

    def extract_estrutura(self, texto_caso, intencao=None):
        self.chamadas.append((texto_caso, intencao))
        return EstruturaArgumentativa(pedido_principal="p", tese_central=texto_caso)


class _Store:
    def __init__(self, teses=()):
        self.teses = list(teses)

    def listar_teses(self, split=None):
        return list(self.teses)

    def criar_tese(self, estrutura, origem, tema, split, documento_origem_id=None):
        t = Tese(uuid.uuid4(), estrutura, origem, tema, split, documento_origem_id)
        self.teses.append(t)
        return t


@pytest.fixture
def acordaos(tmp_path, monkeypatch):
    caminho = tmp_path / "acordaos.json"
    caminho.write_text(json.dumps([{"n": 1, "id_documento": "A1"}, {"n": 2, "id_documento": "A2"}]))
    monkeypatch.setattr(ta, "uuids_origem", lambda url, docs: {"A1": "uuid-1", "A2": "uuid-2"})
    return caminho


def test_ler_planilha(tmp_path):
    p = _planilha(tmp_path / "p.xlsx", [
        (1, "periculosidade", "  texto um ", LADO_RECLAMANTE),
        (2, "horas extras", None, None),
        (3, None, "texto três", LADO_RECLAMADA),
    ])
    linhas = ta.ler_planilha_teses(p)
    assert linhas == [
        {"n": 1, "tema": "periculosidade", "texto": "texto um", "lado": LADO_RECLAMANTE, "email": "adv@x.com"},
        {"n": 3, "tema": None, "texto": "texto três", "lado": LADO_RECLAMADA, "email": "adv@x.com"},
    ]


def test_ler_planilha_sem_email_falha(tmp_path):
    p = _planilha(tmp_path / "p.xlsx", [(1, "x", "texto", LADO_RECLAMANTE)], email=None)
    with pytest.raises(ValueError, match="e-mail"):
        ta.ler_planilha_teses(p)


def test_ler_planilha_vazia_sem_email_nao_falha(tmp_path):
    p = _planilha(tmp_path / "p.xlsx", [(1, "x", None, None)], email=None)
    assert ta.ler_planilha_teses(p) == []


def test_importar_cria_tese_do_advogado(tmp_path, acordaos):
    p = _planilha(tmp_path / "p.xlsx", [
        (1, "periculosidade", "texto um", LADO_RECLAMANTE),
        (2, "horas extras", "texto dois", LADO_RECLAMADA),
    ])
    store, enr = _Store(), _Enricher()
    assert ta.importar(store, enr, p, "db", 42, acordaos_json=acordaos) == 2
    assert enr.chamadas == [("texto um", LADO_RECLAMANTE), ("texto dois", LADO_RECLAMADA)]
    t1, t2 = store.teses
    assert (t1.origem, t1.documento_origem_id, t1.tema) == ("advogado", "uuid-1", "periculosidade")
    assert t1.estrutura["_texto_original"] == "texto um"
    assert t1.estrutura["_avaliador"] == "adv@x.com"
    assert t1.estrutura["_lado"] == "reclamante"
    assert t2.estrutura["_lado"] == "reclamada" and t2.documento_origem_id == "uuid-2"
    from avaliacao.geracao import sortear_split
    assert t1.split == sortear_split(TEMAS, 42)["periculosidade"]


def test_importar_idempotente(tmp_path, acordaos):
    p = _planilha(tmp_path / "p.xlsx", [(1, "periculosidade", "texto um", LADO_RECLAMANTE)])
    store = _Store()
    assert ta.importar(store, _Enricher(), p, "db", 42, acordaos_json=acordaos) == 1
    enr = _Enricher()
    assert ta.importar(store, enr, p, "db", 42, acordaos_json=acordaos) == 0
    assert enr.chamadas == [] and len(store.teses) == 1


def test_importar_tema_vazio_usa_atribuir_tema(tmp_path, acordaos, monkeypatch):
    p = _planilha(tmp_path / "p.xlsx", [(1, None, "texto um", LADO_RECLAMANTE)])
    monkeypatch.setattr(ta, "atribuir_tema", lambda estrutura, client=None: "insalubridade")
    store = _Store()
    ta.importar(store, _Enricher(), p, "db", 42, acordaos_json=acordaos)
    assert store.teses[0].tema == "insalubridade"


def test_importar_lado_desconhecido_falha(tmp_path, acordaos):
    p = _planilha(tmp_path / "p.xlsx", [(1, "periculosidade", "texto", "sei lá")])
    with pytest.raises(ValueError, match="lado"):
        ta.importar(_Store(), _Enricher(), p, "db", 42, acordaos_json=acordaos)


def test_cli_conta_importadas_e_puladas(tmp_path, acordaos, monkeypatch, capsys):
    p = _planilha(tmp_path / "p.xlsx", [(1, "periculosidade", "texto um", LADO_RECLAMANTE)])
    store = _Store()
    monkeypatch.setattr(ta, "_dependencias", lambda: (store, _Enricher(), "db"))
    monkeypatch.setattr(ta, "ACORDAOS_JSON", acordaos)
    assert ta.main(["importar", str(p)]) == 0
    assert ta.main(["importar", str(p)]) == 0
    saida = capsys.readouterr().out
    assert "1 importada(s), 0 pulada(s)" in saida
    assert "0 importada(s), 1 pulada(s)" in saida


def test_cli_retorna_1_se_arquivo_falha(tmp_path, monkeypatch):
    monkeypatch.setattr(ta, "_dependencias", lambda: (_Store(), _Enricher(), "db"))
    assert ta.main(["importar", str(tmp_path / "nao-existe.xlsx")]) == 1


DATABASE_URL = os.getenv("DATABASE_URL")


def _pg():
    if not DATABASE_URL:
        return False
    try:
        psycopg.connect(DATABASE_URL, connect_timeout=2).close()
        return True
    except psycopg.OperationalError:
        return False


@pytest.mark.skipif(not _pg(), reason="Requer Postgres real (DATABASE_URL).")
def test_ensure_schema_aceita_origem_advogado_duas_vezes():
    from avaliacao.db import ensure_schema
    from avaliacao.store import AvaliacaoStore

    ensure_schema(DATABASE_URL)
    ensure_schema(DATABASE_URL)
    store = AvaliacaoStore(DATABASE_URL)
    tese = store.criar_tese({"_texto_original": "x"}, "advogado", "outro", "dev")
    try:
        assert tese.origem == "advogado"
    finally:
        with psycopg.connect(DATABASE_URL) as conn:
            conn.execute("DELETE FROM avaliacao.teses WHERE id = %s", (tese.id,))


def test_ler_planilha_real_email_na_coluna_c(tmp_path):
    """A planilha distribuída tem rótulos em A:B mesclado e respostas em C."""
    import shutil
    from pathlib import Path

    origem = Path(ta.__file__).resolve().parent.parent / "docs" / "avaliacao" / "1-producao-de-teses.xlsx"
    p = tmp_path / "real.xlsx"
    shutil.copy(origem, p)
    wb = openpyxl.load_workbook(p)
    ws = wb["Teses"]
    ws["C4"] = "Adv@X.com"
    assert ws["A11"].value == 2
    ws["E11"], ws["F11"] = "tese da linha 11", LADO_RECLAMADA
    wb.save(p)
    linhas = ta.ler_planilha_teses(p)
    assert linhas == [{"n": 2, "tema": "insalubridade", "texto": "tese da linha 11",
                       "lado": LADO_RECLAMADA, "email": "Adv@X.com"}]


def test_importar_tema_vazio_nao_manda_chaves_sublinhado(tmp_path, acordaos, monkeypatch):
    p = _planilha(tmp_path / "p.xlsx", [(1, None, "texto um", LADO_RECLAMANTE)])
    recebido = []
    monkeypatch.setattr(ta, "atribuir_tema", lambda estrutura, client=None: recebido.append(estrutura) or "outro")
    ta.importar(_Store(), _Enricher(), p, "db", 42, acordaos_json=acordaos)
    assert recebido and not any(k.startswith("_") for k in recebido[0])


def test_importar_idempotente_com_email_em_caixa_diferente(tmp_path, acordaos):
    store = _Store()
    p1 = _planilha(tmp_path / "a.xlsx", [(1, "periculosidade", "texto um", LADO_RECLAMANTE)], email="Adv@X.com")
    p2 = _planilha(tmp_path / "b.xlsx", [(1, "periculosidade", "texto um", LADO_RECLAMANTE)], email="adv@x.com")
    assert ta.importar(store, _Enricher(), p1, "db", 42, acordaos_json=acordaos) == 1
    assert ta.importar(store, _Enricher(), p2, "db", 42, acordaos_json=acordaos) == 0


def test_dependencias_carregam_o_env(monkeypatch):
    import avaliacao.contexto as ctx

    chamadas = []
    monkeypatch.setattr(ctx, "_database_url", lambda: chamadas.append(1) or "postgresql://x")
    import avaliacao.db as db
    import avaliacao.store as st
    import query.enrichment.query_enrichment as qe
    monkeypatch.setattr(db, "ensure_schema", lambda url: None)
    monkeypatch.setattr(qe, "QueryEnricher", lambda: "enr")
    monkeypatch.setattr(st, "AvaliacaoStore", lambda url: ("store", url))
    assert ta._dependencias() == (("store", "postgresql://x"), "enr", "postgresql://x")
    assert chamadas
