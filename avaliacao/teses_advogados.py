"""
Teses escritas por advogados: importa a planilha de produção de teses preenchida
para o conjunto-ouro (origem "advogado"). O texto vira estrutura argumentativa pelo
mesmo caminho do /enrich (QueryEnricher) e o texto original fica guardado, para
servir de exemplo de estilo na geração das sintéticas.

    python -m avaliacao.teses_advogados importar ARQUIVO...
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from types import SimpleNamespace

import openpyxl

from avaliacao.geracao import sortear_split, uuids_origem
from avaliacao.temas import TEMAS, atribuir_tema

ACORDAOS_JSON = Path(__file__).resolve().parent.parent / "docs" / "avaliacao" / "1-acordaos-selecionados.json"

_ROTULOS_DADOS = {"e-mail": "email"}


def ler_planilha_teses(caminho) -> list[dict]:
    """Linhas preenchidas da aba "Teses": n, tema, texto, lado (texto da lista) e email.

    A tabela é localizada pelo cabeçalho "Nº"; linhas sem tese (e o exemplo "Ex.") são
    ignoradas. Linha com tese e sem e-mail do avaliador gera ValueError."""
    ws = openpyxl.load_workbook(caminho, data_only=True)["Teses"]
    email = None
    cabecalho = None
    for linha in ws.iter_rows(min_col=1, max_col=7):
        rotulo = str(linha[0].value or "").strip()
        if rotulo.lower() in _ROTULOS_DADOS and cabecalho is None:
            # rótulo em A:B mesclado → resposta em C; rótulo só em A → resposta em B
            email = next((str(c.value).strip() for c in linha[1:3] if str(c.value or "").strip()), None)
        elif rotulo == "Nº":
            cabecalho = linha[0].row
            break
    if cabecalho is None:
        raise ValueError(f"{caminho}: cabeçalho 'Nº' não encontrado na aba Teses")

    saida = []
    for linha in ws.iter_rows(min_row=cabecalho + 1, max_col=7):
        n = linha[0].value
        texto = str(linha[4].value or "").strip()
        if not isinstance(n, int) or not texto:
            continue
        if not email:
            raise ValueError(f"{caminho}: há tese preenchida, mas falta o e-mail do avaliador")
        saida.append({
            "n": n,
            "tema": str(linha[1].value or "").strip() or None,
            "texto": texto,
            "lado": str(linha[5].value or "").strip() or None,
            "email": email,
        })
    return saida


def _lado(texto: str | None) -> str | None:
    if texto is None:
        return None
    baixo = texto.lower()
    if "reclamante" in baixo:
        return "reclamante"
    if "reclamada" in baixo:
        return "reclamada"
    raise ValueError(f"lado desconhecido: {texto!r}")


def _importar(store, enricher, caminho, database_url, seed, acordaos_json=None) -> tuple[int, int]:
    """(importadas, puladas por já existirem)."""
    linhas = ler_planilha_teses(caminho)
    if not linhas:
        return 0, 0
    acordaos = json.loads(Path(acordaos_json or ACORDAOS_JSON).read_text(encoding="utf-8"))
    por_n = {a["n"]: a["id_documento"] for a in acordaos}
    ids = {l["n"]: por_n[l["n"]] for l in linhas if l["n"] in por_n}
    faltando = [l["n"] for l in linhas if l["n"] not in por_n]
    if faltando:
        raise ValueError(f"{caminho}: Nº sem acórdão em {ACORDAOS_JSON.name}: {faltando}")
    uuids = uuids_origem(database_url, [SimpleNamespace(id_documento=i) for i in ids.values()])

    ja = {
        (t.documento_origem_id, (t.estrutura.get("_avaliador") or "").lower(),
         (t.estrutura.get("_texto_original") or "").strip())
        for t in store.listar_teses() if t.origem == "advogado"
    }
    splits = sortear_split(TEMAS, seed)
    importadas = puladas = 0
    for l in linhas:
        uid = uuids.get(ids[l["n"]])
        if uid is None:
            raise ValueError(f"{caminho}: acórdão {ids[l['n']]} (Nº {l['n']}) não existe em documentos")
        lado = _lado(l["lado"])
        chave = (uid, l["email"].lower(), l["texto"])
        if chave in ja:
            puladas += 1
            continue
        estrutura = enricher.extract_estrutura(l["texto"], intencao=l["lado"]).model_dump()
        estrutura.update(_texto_original=l["texto"], _avaliador=l["email"], _lado=lado)
        tema = l["tema"] if l["tema"] in splits else atribuir_tema({k: v for k, v in estrutura.items() if not k.startswith("_")})
        store.criar_tese(estrutura, "advogado", tema, splits[tema], documento_origem_id=uid)
        ja.add(chave)
        importadas += 1
    return importadas, puladas


def importar(store, enricher, caminho, database_url, seed, acordaos_json=None) -> int:
    """Importa as teses da planilha; devolve quantas criou (as já existentes são puladas)."""
    return _importar(store, enricher, caminho, database_url, seed, acordaos_json)[0]


def _dependencias():
    from avaliacao.contexto import _database_url  # carrega o .env
    from avaliacao.db import ensure_schema
    from avaliacao.store import AvaliacaoStore
    from query.enrichment.query_enrichment import QueryEnricher

    url = _database_url()
    ensure_schema(url)
    return AvaliacaoStore(url), QueryEnricher(), url


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python -m avaliacao.teses_advogados", description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="comando", required=True)
    imp = sub.add_parser("importar", help="importa planilhas de teses preenchidas")
    imp.add_argument("arquivos", nargs="+")
    args = parser.parse_args(argv)

    import os

    store, enricher, url = _dependencias()  # antes do seed: carrega o .env
    seed = int(os.getenv("AVALIACAO_SEED", "42"))
    falhas = 0
    for arquivo in args.arquivos:
        try:
            novas, puladas = _importar(store, enricher, arquivo, url, seed, ACORDAOS_JSON)
        except Exception as e:  # noqa: BLE001 — um arquivo ruim não impede os outros
            print(f"{arquivo}: falhou: {e}", file=sys.stderr)
            falhas += 1
            continue
        print(f"{arquivo}: {novas} importada(s), {puladas} pulada(s)")
    return 1 if falhas else 0


if __name__ == "__main__":
    sys.exit(main())
