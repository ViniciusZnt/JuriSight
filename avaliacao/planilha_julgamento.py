"""Planilha de julgamento: exporta um lote por tese e importa o que voltou.

Uso:
    python -m avaliacao.planilha_julgamento exportar [--teses id,id | --todas] [--pasta data/eval/lotes] [--seed 42]
    python -m avaliacao.planilha_julgamento importar ARQUIVO [ARQUIVO ...]

A tese do lote fica no nome definido `tese_id` do .xlsx (constante de texto,
sem célula nem aba extra); o `par_id` de cada linha fica na coluna A, oculta.
A proteção da planilha não tem senha: só evita edição acidental fora das
células amarelas.
"""
from __future__ import annotations

import argparse
import os
import random
import re
import sys
import unicodedata
import uuid
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

STANCES = ("fortalece", "enfraquece", "neutro")
ATUACOES = ("Reclamante", "Reclamada", "Ambos", "Outra")
NOME_TESE_ID = "tese_id"
CORREDORES_DO_POOL = ("hibrido", "bm25", "vetorial")

COL_PAR, COL_NUM, COL_PROCESSO, COL_LINK, COL_EMENTA, COL_FUND, COL_NOTA, COL_STANCE, COL_COMENT = range(1, 10)
LINHA_CABECALHO = 8
LINHA_INICIAL = LINHA_CABECALHO + 1
CHARS_DISPOSITIVO = 1500
CHARS_FUNDAMENTACAO = 3000

AMARELO = PatternFill("solid", fgColor="FFF2CC")
CINZA = PatternFill("solid", fgColor="F2F2F2")
ESCURO = PatternFill("solid", fgColor="3B3B3B")
_QUEBRA = Alignment(wrap_text=True, vertical="top")

_RUBRICA = (
    ("0", "Outro problema jurídico."),
    ("1", "Mesmo ramo ou tema, mas outro pedido ou outro fato central."),
    ("2", "Mesmo pedido, mas o fato decisivo é diferente ou o acórdão cobre só uma parte."),
    ("3", "Mesmo problema e mesmo fato decisivo: você citaria este acórdão na peça."),
)
_INSTRUCOES = (
    ("Julgamento de acórdãos — benchmark JuriSight", "titulo"),
    ("", None),
    ("Para que serve", "negrito"),
    ("Cada linha da aba Avaliação é um acórdão. Diga o quanto ele ajuda a tese descrita no topo da aba.", None),
    ("", None),
    ("Como preencher", "negrito"),
    ("1. Na aba Avaliador, preencha nome, e-mail e atuação (células amarelas). O e-mail é obrigatório.", None),
    ("2. Leia a aba Rubrica: ela define a nota de 0 a 3.", None),
    ("3. Na aba Avaliação, leia a tese e, para cada acórdão, a ementa (ou o dispositivo) e o início da fundamentação.", None),
    ("4. Escolha a Nota (0 a 3) na lista.", None),
    ("5. Com nota 2 ou 3, escolha também o Stance: fortalece (o acórdão apoia a tese), enfraquece (vai contra) ou neutro. Com nota 0 ou 1 deixe em branco.", None),
    ("6. O comentário é opcional. Não apague nem reordene linhas; a coluna A (oculta) identifica o acórdão.", None),
    ("7. Faça o julgamento sozinho, sem combinar com o outro avaliador. Devolva o arquivo .xlsx.", None),
)


def _slug(texto: str) -> str:
    base = unicodedata.normalize("NFKD", texto or "").encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", base.lower()).strip("-") or "sem-tema"


def nome_arquivo(tese, dia: str | None = None) -> str:
    return f"lote_{dia or date.today().isoformat()}_{_slug(tese.tema)}_{str(tese.id)[:8]}.xlsx"


def _aba_instrucoes(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Instruções"
    ws.column_dimensions["A"].width = 110
    for i, (texto, estilo) in enumerate(_INSTRUCOES, start=1):
        c = ws.cell(i, 1, texto)
        c.alignment = _QUEBRA
        c.font = Font(size=14 if estilo == "titulo" else 10, bold=estilo in ("titulo", "negrito"))


def _aba_rubrica(wb: Workbook) -> None:
    ws = wb.create_sheet("Rubrica")
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 100
    for i, (nota, texto) in enumerate(_RUBRICA, start=1):
        ws.cell(i, 1, nota).font = Font(bold=True, size=10)
        ws.cell(i, 2, texto).alignment = _QUEBRA


def _aba_avaliador(wb: Workbook) -> None:
    ws = wb.create_sheet("Avaliador")
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 40
    for i, rotulo in enumerate(("Nome", "E-mail", "Atuação"), start=1):
        c = ws.cell(i, 1, rotulo)
        c.fill, c.font = CINZA, Font(bold=True, size=10)
        r = ws.cell(i, 2)
        r.fill, r.protection = AMARELO, Protection(locked=False)
    dv = DataValidation(type="list", formula1='"' + ",".join(ATUACOES) + '"', allow_blank=True,
                        showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add("B3")
    ws.protection.sheet = True


def _texto_da_tese(estrutura: dict) -> list[tuple[str, str]]:
    normas = estrutura.get("normas") or []
    if not isinstance(normas, str):
        normas = "; ".join(str(n) for n in normas)
    linhas = [
        ("Pedido principal", str(estrutura.get("pedido_principal") or "")),
        ("Tese central", str(estrutura.get("tese_central") or "")),
        ("Normas", normas),
    ]
    if estrutura.get("_texto_original"):
        linhas.append(("Texto original", str(estrutura["_texto_original"])))
    return linhas


def _trecho_ementa(doc) -> str:
    if doc.ementa and doc.ementa.strip():
        return doc.ementa.strip()
    return (doc.acordao or "").strip()[-CHARS_DISPOSITIVO:]


def exportar_lote(tese, pares: list, documentos: dict, caminho: Path, seed: int) -> int:
    """Grava o .xlsx do lote da tese e devolve o número de linhas de pares."""
    wb = Workbook()
    _aba_instrucoes(wb)
    _aba_rubrica(wb)
    _aba_avaliador(wb)
    ws = wb.create_sheet("Avaliação")
    wb.defined_names[NOME_TESE_ID] = DefinedName(NOME_TESE_ID, attr_text=f'"{tese.id}"')

    larguras = {COL_PAR: 38, COL_NUM: 5, COL_PROCESSO: 24, COL_LINK: 24, COL_EMENTA: 70,
                COL_FUND: 70, COL_NOTA: 8, COL_STANCE: 14, COL_COMENT: 36}
    for col, w in larguras.items():
        ws.column_dimensions[ws.cell(1, col).column_letter].width = w
    ws.column_dimensions["A"].hidden = True

    ws.cell(1, COL_NUM, "Tese a julgar").font = Font(bold=True, size=14)
    for i, (rotulo, texto) in enumerate(_texto_da_tese(tese.estrutura), start=2):
        r = ws.cell(i, COL_NUM, rotulo)
        r.font, r.alignment = Font(bold=True, size=10), _QUEBRA
        ws.cell(i, COL_PROCESSO, texto).alignment = _QUEBRA
        ws.merge_cells(start_row=i, start_column=COL_PROCESSO, end_row=i, end_column=COL_FUND)
        ws.row_dimensions[i].height = 60 if rotulo == "Texto original" else 32

    cabecalho = {COL_PAR: "Código", COL_NUM: "Nº", COL_PROCESSO: "Processo", COL_LINK: "Link",
                 COL_EMENTA: "Ementa (ou dispositivo)", COL_FUND: "Início da fundamentação",
                 COL_NOTA: "Nota", COL_STANCE: "Stance", COL_COMENT: "Comentário"}
    for col, texto in cabecalho.items():
        c = ws.cell(LINHA_CABECALHO, col, texto)
        c.fill, c.font, c.alignment = ESCURO, Font(bold=True, color="FFFFFF", size=10), _QUEBRA

    ordem = sorted(pares, key=lambda p: str(p.id))
    random.Random(f"{seed}:{tese.id}").shuffle(ordem)
    for n, par in enumerate(ordem, start=1):
        doc = documentos.get(par.documento_id)
        linha = LINHA_INICIAL + n - 1
        valores = {
            COL_PAR: str(par.id), COL_NUM: n,
            COL_PROCESSO: getattr(doc, "numero_processo", "") or "",
            COL_LINK: getattr(doc, "link_original", "") or "",
            COL_EMENTA: _trecho_ementa(doc) if doc else "(documento indisponível)",
            COL_FUND: ((doc.fundamentacao or "")[:CHARS_FUNDAMENTACAO]) if doc else "",
        }
        for col in range(1, COL_COMENT + 1):
            c = ws.cell(linha, col, valores.get(col))
            c.alignment = _QUEBRA
            c.font = Font(size=10)
            if col >= COL_NOTA:
                c.fill, c.protection = AMARELO, Protection(locked=False)
        ws.row_dimensions[linha].height = 150

    ate = LINHA_INICIAL + max(len(ordem), 1) - 1
    dv_nota = DataValidation(type="list", formula1='"0,1,2,3"', allow_blank=True,
                             showErrorMessage=True, errorTitle="Nota", error="Escolha de 0 a 3.")
    dv_stance = DataValidation(type="list", formula1='"' + ",".join(STANCES) + '"', allow_blank=True,
                               showErrorMessage=True, errorTitle="Stance",
                               error="Escolha fortalece, enfraquece ou neutro.")
    for dv, col in ((dv_nota, "G"), (dv_stance, "H")):
        ws.add_data_validation(dv)
        dv.add(f"{col}{LINHA_INICIAL}:{col}{ate}")
    ws.freeze_panes = f"A{LINHA_INICIAL}"
    # formatColumns/Rows liberados para o avaliador poder redimensionar
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False

    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
    return len(ordem)


# --- leitura ---

def _vazio(v) -> bool:
    return v is None or str(v).strip() == ""


def _nota(v) -> int | None:
    try:
        f = float(str(v).strip().replace(",", "."))
    except ValueError:
        return None
    return int(f) if f in (0.0, 1.0, 2.0, 3.0) else None


def _ler(caminho):
    """(email, tese_id, validas, puladas[(linha, motivo)]) de um lote."""
    wb = load_workbook(caminho, data_only=True)
    try:
        email = str(wb["Avaliador"]["B2"].value or "").strip().lower()
        ws = wb["Avaliação"]
    except KeyError as e:
        raise ValueError(f"{Path(caminho).name}: não parece um lote de julgamento (falta a aba {e}).") from e
    if not email or "@" not in email:
        raise ValueError(f"{Path(caminho).name}: e-mail do avaliador ausente ou inválido (aba Avaliador, célula B2).")
    nome = wb.defined_names.get(NOME_TESE_ID)
    try:
        tese_id = uuid.UUID((nome.attr_text or "").strip('"'))
    except (AttributeError, ValueError) as e:
        raise ValueError(f"{Path(caminho).name}: o arquivo não identifica a tese do lote.") from e

    validas, puladas, vistos, com_nota = [], [], set(), 0
    for row in ws.iter_rows(min_row=LINHA_INICIAL):
        linha = row[0].row
        bruto_par, bruto_nota, bruto_stance = (row[c - 1].value for c in (COL_PAR, COL_NOTA, COL_STANCE))
        if _vazio(bruto_par) and _vazio(bruto_nota):
            continue
        if _vazio(bruto_nota):
            puladas.append((linha, "sem nota"))
            continue
        com_nota += 1
        try:
            par_id = uuid.UUID(str(bruto_par).strip())
        except ValueError:
            puladas.append((linha, "código do par ausente ou alterado"))
            continue
        nota = _nota(bruto_nota)
        if nota is None:
            puladas.append((linha, "nota fora de 0 a 3"))
            continue
        stance = "" if _vazio(bruto_stance) else str(bruto_stance).strip().lower()
        if stance and stance not in STANCES:
            puladas.append((linha, "stance inválido"))
            continue
        if nota >= 2 and not stance:
            puladas.append((linha, "stance vazio com nota 2 ou 3"))
            continue
        if par_id in vistos:
            puladas.append((linha, "par repetido"))
            continue
        vistos.add(par_id)
        validas.append((par_id, nota, stance if nota >= 2 else "neutro"))
    if not com_nota:
        raise ValueError(f"{Path(caminho).name}: nenhuma nota preenchida.")
    return email, tese_id, validas, puladas


def _avisar(caminho, puladas) -> None:
    for linha, motivo in puladas:
        if motivo != "sem nota":
            print(f"{Path(caminho).name}: linha {linha} ignorada ({motivo})", file=sys.stderr)


def ler_lote(caminho) -> tuple[str, list[tuple[uuid.UUID, int, str]]]:
    """(e-mail normalizado, linhas válidas (par_id, nota, stance)). Linhas
    inválidas são avisadas em stderr; e-mail ausente ou sem nota → ValueError."""
    email, _, validas, puladas = _ler(caminho)
    _avisar(caminho, puladas)
    return email, validas


def importar_arquivo(store, caminho) -> tuple[int, dict[str, int]]:
    """Grava os julgamentos do arquivo. Devolve (gravadas, {motivo: puladas}).
    Idempotente: o store faz upsert por (par, juiz, rodada)."""
    from avaliacao.store import Julgamento

    email, tese_id, validas, puladas = _ler(caminho)
    _avisar(caminho, puladas)
    if not any(t.id == tese_id for t in store.listar_teses()):
        raise ValueError(f"{Path(caminho).name}: a tese {tese_id} não existe no banco.")
    pares_da_tese = {p.id for p in store.pares_da_tese(tese_id)}
    motivos = Counter(m for _, m in puladas)
    gravadas = 0
    for par_id, nota, stance in validas:
        if par_id not in pares_da_tese:
            motivos["par fora da tese"] += 1
            print(f"{Path(caminho).name}: par {par_id} não pertence à tese do lote; ignorado", file=sys.stderr)
            continue
        store.gravar_julgamento(Julgamento(par_id, f"humano:{email}", 1, nota, stance))
        gravadas += 1
    return gravadas, dict(motivos)


# --- CLI ---

def _exportar(args, store) -> int:
    from query.search.doc_retriever import DocRetriever

    teses = store.listar_teses()
    if not args.todas:
        ids = {i.strip() for i in args.teses.split(",") if i.strip()}
        teses = [t for t in teses if str(t.id) in ids]
        faltam = ids - {str(t.id) for t in teses}
        if faltam:
            print(f"teses não encontradas: {', '.join(sorted(faltam))}", file=sys.stderr)
            return 1
    retriever = DocRetriever(os.environ["DATABASE_URL"])
    for tese in teses:
        # pares antigos (pipeline, passo_zero) ficam fora do lote
        pares = [p for p in store.pares_da_tese(tese.id) if set(p.corredores) & set(CORREDORES_DO_POOL)]
        docs = retriever.fetch([p.documento_id for p in pares])
        destino = Path(args.pasta) / nome_arquivo(tese)
        n = exportar_lote(tese, pares, docs, destino, args.seed)
        print(f"{destino}: {n} pares")
    return 0


def main(argv=None, store=None) -> int:
    parser = argparse.ArgumentParser(prog="python -m avaliacao.planilha_julgamento", description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="cmd", required=True)
    exp = sub.add_parser("exportar", help="gera um .xlsx por tese")
    alvo = exp.add_mutually_exclusive_group(required=True)
    alvo.add_argument("--teses", help="ids das teses, separados por vírgula")
    alvo.add_argument("--todas", action="store_true")
    exp.add_argument("--pasta", default="data/eval/lotes")
    exp.add_argument("--seed", type=int, default=42, help="semente da ordem das linhas (padrão 42)")
    imp = sub.add_parser("importar", help="grava os julgamentos dos arquivos devolvidos")
    imp.add_argument("arquivos", nargs="+", metavar="ARQUIVO")
    args = parser.parse_args(argv)

    if store is None:
        from avaliacao.contexto import carregar_store
        store = carregar_store()
    if args.cmd == "exportar":
        return _exportar(args, store)

    falhou = False
    for arquivo in args.arquivos:
        try:
            gravadas, puladas = importar_arquivo(store, arquivo)
        except (ValueError, OSError) as e:
            print(f"{arquivo}: FALHOU ({e})", file=sys.stderr)
            falhou = True
            continue
        extra = f"; puladas {sum(puladas.values())} ({', '.join(f'{m}: {n}' for m, n in puladas.items())})" if puladas else "; puladas 0"
        print(f"{arquivo}: {gravadas} gravada(s){extra}")
    return 1 if falhou else 0


if __name__ == "__main__":
    sys.exit(main())
